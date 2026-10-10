"""Turn outreach-contacts.csv into Outreach-Contacts.xlsx in the Outbound Engine folder, one tab per vertical."""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), "Outreach-Contacts.xlsx")
RED = "FF3838"
COLS = [("#", "row", 5), ("Company", "company", 26), ("Contact", "contact_name", 24),
        ("Role", "contact_role", 30), ("LinkedIn profile", "linkedin_profile", 34),
        ("Second contact", "second_contact", 30), ("Second contact LinkedIn", "second_contact_linkedin", 30),
        ("Company LinkedIn", "company_linkedin", 24), ("Verified email (general)", "verified_general_email", 24),
        ("Personal email (add after Hunter/Apollo check)", "personal_email", 30),
        ("Status", "status", 16), ("Note", "note", 60)]

rows = list(csv.DictReader(open(os.path.join(HERE, "outreach-contacts.csv"))))
wb = Workbook()
wb.remove(wb.active)
for vertical in dict.fromkeys(r["vertical"] for r in rows):
    ws = wb.create_sheet(vertical[:31])
    for i, (label, _, width) in enumerate(COLS, 1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=RED)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "C2"
    for r_i, r in enumerate((x for x in rows if x["vertical"] == vertical), 2):
        for c_i, (_, key, _) in enumerate(COLS, 1):
            val = r[key]
            c = ws.cell(row=r_i, column=c_i, value=int(val) if key == "row" else val)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if key.endswith("linkedin") or key == "linkedin_profile":
                if val.startswith("http"):
                    c.hyperlink = val
                    c.font = Font(color=RED, underline="single")
wb.save(OUT)
print("wrote", OUT)
